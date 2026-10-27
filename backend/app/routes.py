import hashlib
import io
import csv
import secrets
import smtplib
import zipfile
from datetime import datetime, timedelta, timezone
from email.message import EmailMessage
from pathlib import Path
from flask import Blueprint, abort, current_app, g, request, session, send_file
from flask_wtf.csrf import generate_csrf
from sqlalchemy import or_
from werkzeug.security import check_password_hash, generate_password_hash
from werkzeug.utils import secure_filename
from .extensions import db, limiter
from .models import *
from .schemas import *
from .services import *

api = Blueprint('api', __name__)


@api.get('/auth/csrf')
def csrf_token():
    return {'csrf_token': generate_csrf()}


@api.post('/auth/login')
@limiter.limit('10 per minute')
def login():
    data = parse(Login)
    user = User.query.filter(or_(User.username == data.username, User.email == data.username.lower())).first()
    # Perform an expensive hash check even when the username does not exist.
    if '_DUMMY_HASH' not in current_app.config:
        current_app.config['_DUMMY_HASH'] = generate_password_hash(secrets.token_hex(16))
    fallback = current_app.config['_DUMMY_HASH']
    valid = check_password_hash(user.password_hash if user else fallback, data.password)
    if not user or not valid or not user.active:
        abort(401, description='Invalid credentials or inactive account.')
    session.clear()
    session.update(uid=user.id, version=user.session_version)
    session.permanent = True
    g.user = user
    audit('user', user.id, 'login')
    db.session.commit()
    return {'user':user_json(user), 'csrf_token':generate_csrf()}


@api.post('/auth/logout')
@authenticated
def logout():
    # Revoke copied session cookies as well as the browser's current cookie.
    g.user.session_version += 1
    audit('user', g.user.id, 'logout')
    db.session.commit()
    session.clear()
    return {'message':'Signed out'}


@api.get('/auth/me')
@authenticated
def me():
    return user_json(g.user)


@api.patch('/auth/profile')
@authenticated
def profile():
    data = request.get_json() or {}
    if not data or set(data)-{'full_name','phone','position'}:
        abort(422, description='Only full_name, phone and position may be changed.')
    current = user_json(g.user)
    current.update(data)
    validated = parse(EmployeeInput, {k:current[k] for k in EmployeeInput.model_fields if k in current})
    for key in data:
        setattr(g.user,key,getattr(validated,key))
    audit('user',g.user.id,'profile_updated')
    db.session.commit()
    return user_json(g.user)


@api.post('/auth/change-password')
@authenticated
def change_password():
    data = request.get_json() or {}
    password = data.get('new_password','')
    if not isinstance(password,str) or not 12 <= len(password) <= 200:
        abort(422, description='Password must have 12–200 characters.')
    if not check_password_hash(g.user.password_hash,str(data.get('current_password',''))):
        abort(403, description='Current password is incorrect.')
    g.user.password_hash = generate_password_hash(password)
    g.user.session_version += 1
    session['version'] = g.user.session_version
    audit('user',g.user.id,'password_changed')
    db.session.commit()
    return {'message':'Password changed'}


def send_reset(email, token):
    if not current_app.config.get('MAIL_SERVER'):
        raise RuntimeError('SMTP is not configured')
    message = EmailMessage()
    message['Subject'] = 'Reset your Workday password'
    message['From'] = current_app.config['MAIL_FROM']
    message['To'] = email
    message.set_content(f"Reset your password within 30 minutes: {current_app.config['FRONTEND_URL']}/reset-password?token={token}")
    with smtplib.SMTP(current_app.config['MAIL_SERVER'],current_app.config['MAIL_PORT'],timeout=15) as smtp:
        if current_app.config['MAIL_USE_TLS']:
            smtp.starttls()
        if current_app.config['MAIL_USERNAME']:
            smtp.login(current_app.config['MAIL_USERNAME'],current_app.config['MAIL_PASSWORD'])
        smtp.send_message(message)


@api.post('/auth/forgot-password')
@limiter.limit('5 per hour')
def forgot_password():
    from pydantic import TypeAdapter, EmailStr, ValidationError
    try:
        email = str(TypeAdapter(EmailStr).validate_python((request.get_json() or {}).get('email'))).lower()
    except ValidationError:
        abort(422, description='Enter a valid email address.')
    user = User.query.filter_by(email=email,active=True).first()
    if user:
        token = secrets.token_urlsafe(32)
        reset = PasswordReset(user_id=user.id,token_hash=hashlib.sha256(token.encode()).hexdigest(),expires_at=now()+timedelta(minutes=30))
        db.session.add(reset)
        try:
            send_reset(email,token)
            db.session.commit()
        except Exception:
            db.session.rollback()
            current_app.logger.error('Password reset email delivery failed; check SMTP configuration.')
    return {'message':'If an active account matches, a reset link will be sent.'}


@api.post('/auth/reset-password')
@limiter.limit('10 per hour')
def reset_password():
    data = request.get_json() or {}
    token = str(data.get('token',''))
    password = data.get('password','')
    if not isinstance(password,str) or not 12 <= len(password) <= 200:
        abort(422, description='Password must have 12–200 characters.')
    reset = PasswordReset.query.filter_by(token_hash=hashlib.sha256(token.encode()).hexdigest(),used=False).with_for_update().first()
    if not reset or reset.expires_at.replace(tzinfo=timezone.utc) < now():
        abort(422, description='Reset link is invalid or expired.')
    user = db.session.get(User,reset.user_id)
    if not user.active:
        abort(422, description='Reset link is invalid or expired.')
    user.password_hash = generate_password_hash(password)
    user.session_version += 1
    for item in PasswordReset.query.filter_by(user_id=user.id,used=False):
        item.used = True
    g.user = user
    audit('user',user.id,'password_reset')
    db.session.commit()
    session.clear()
    return {'message':'Password reset. Please sign in.'}


@api.get('/employees')
@authenticated
def employees():
    query = scoped_users()
    if request.args.get('q'):
        term = '%' + request.args['q'][:200] + '%'
        query = query.filter(or_(User.full_name.ilike(term),User.employee_id.ilike(term),User.email.ilike(term)))
    if request.args.get('department'):
        query = query.filter(User.department_id == filter_value('department', User.department_id))
    if request.args.get('role'):
        query = query.filter(User.role == request.args['role'])
    if request.args.get('active'):
        query = query.filter(User.active == (request.args['active']=='true'))
    return paginate(query,user_json,{'newest':User.id.desc(),'name':User.full_name.asc()})


def validate_employee(data, existing=None):
    require_super()
    if data.department_id:
        db.get_or_404(Department,data.department_id)
    if data.manager_id:
        manager = db.get_or_404(User,data.manager_id)
        if manager.role not in ('manager','admin','super_admin') or not manager.active or (existing and manager.id == existing.id):
            abort(422, description='Manager must be an active manager or administrator.')
    if existing and existing.id == g.user.id and (not data.active or data.role != 'super_admin'):
        abort(409, description='You cannot deactivate or demote your own super admin account.')
    if existing and existing.role in ('manager','admin','super_admin') and data.role in STAFF and User.query.filter_by(manager_id=existing.id).first():
        abort(409, description='Reassign this manager’s employees before changing their role.')


@api.post('/employees')
@authenticated
def create_employee():
    data = parse(EmployeeInput)
    validate_employee(data)
    if not data.password:
        abort(422, description='An initial password is required.')
    user = User(**data.model_dump(exclude={'password'}),password_hash=generate_password_hash(data.password))
    db.session.add(user)
    db.session.flush()
    audit('user',user.id,'created')
    db.session.commit()
    return user_json(user),201


@api.get('/employees/<int:id>')
@authenticated
def employee_detail(id):
    return user_json(scoped_users().filter_by(id=id).first_or_404())


@api.patch('/employees/<int:id>')
@authenticated
def update_employee(id):
    require_super()
    user = db.get_or_404(User,id)
    merged = {k:getattr(user,k) for k in EmployeeInput.model_fields if hasattr(user,k)}
    merged.update(request.get_json() or {})
    data = parse(EmployeeInput,merged)
    validate_employee(data,user)
    for key,value in data.model_dump(exclude={'password'}).items():
        setattr(user,key,value)
    if data.password:
        user.password_hash = generate_password_hash(data.password)
    user.session_version += 1
    if user.id == g.user.id:
        session['version'] = user.session_version
    audit('user',id,'updated')
    db.session.commit()
    return user_json(user)


@api.patch('/employees/<int:id>/status')
@authenticated
def employee_status(id):
    require_super()
    data = request.get_json() or {}
    if set(data) != {'active'} or not isinstance(data['active'],bool):
        abort(422, description='Provide active as a boolean.')
    return update_employee(id)


@api.route('/departments',methods=['GET','POST'])
@authenticated
def departments():
    if request.method == 'POST':
        require_super()
        data = parse(NamedInput)
        item = Department(**data.model_dump())
        db.session.add(item)
        db.session.flush()
        audit('department',item.id,'created')
        db.session.commit()
        return {'id':item.id,'name':item.name,'description':item.description},201
    return {'items':[{'id':d.id,'name':d.name,'description':d.description} for d in Department.query.order_by(Department.name)]}


@api.patch('/departments/<int:id>')
@authenticated
def update_department(id):
    require_super()
    item = db.get_or_404(Department,id)
    data = parse(NamedInput,{'name':item.name,'description':item.description}|(request.get_json() or {}))
    item.name,item.description = data.name,data.description
    audit('department',id,'updated')
    db.session.commit()
    return {'id':id,'name':item.name,'description':item.description}


@api.route('/categories',methods=['GET','POST'])
@authenticated
def categories():
    if request.method == 'POST':
        require_super()
        data = parse(NamedInput)
        item = Category(name=data.name)
        db.session.add(item)
        db.session.flush()
        audit('category',item.id,'created')
        db.session.commit()
        return {'id':item.id,'name':item.name},201
    return {'items':[{'id':c.id,'name':c.name} for c in Category.query.order_by(Category.name)]}


@api.patch('/categories/<int:id>')
@authenticated
def update_category(id):
    require_super()
    item = db.get_or_404(Category,id)
    data = parse(NamedInput)
    item.name = data.name
    audit('category',id,'updated')
    db.session.commit()
    return {'id':id,'name':item.name}


def set_assignees(task, ids):
    users = [db.get_or_404(User,id) for id in ids]
    if not users or len(set(ids)) != len(ids):
        abort(422, description='Select unique assignees.')
    for u in users:
        if not u.active or not (can_manage(u) or u.id == g.user.id):
            abort(403, description='Assignees must be active members of your team.')
    previous = {a.user_id for a in task.assignments}
    for a in list(task.assignments):
        if a.user_id not in ids:
            task.assignments.remove(a)
            notify(a.user_id,f'You were unassigned from {task.title}.','/tasks')
    for u in users:
        if u.id not in previous:
            task.assignments.append(TaskAssignment(user_id=u.id))
            notify(u.id,f'New task assigned: {task.title}',f'/tasks/{task.id}')


def apply_task(task,data):
    if data.department_id:
        db.get_or_404(Department,data.department_id)
    if data.category_id:
        db.get_or_404(Category,data.category_id)
    for key,value in data.model_dump(exclude={'assigned_employee_ids'}).items():
        setattr(task,key,value)


@api.get('/tasks')
@authenticated
def tasks():
    return paginate(task_filters(scoped_tasks()),task_json,{'newest':Task.id.desc(),'deadline':Task.deadline.asc(),'title':Task.title.asc()})


@api.post('/tasks')
@authenticated
def create_task():
    require_manager()
    data = parse(TaskInput)
    task = Task(creator_id=g.user.id)
    apply_task(task,data)
    db.session.add(task)
    db.session.flush()
    set_assignees(task,data.assigned_employee_ids)
    audit('task',task.id,'created',{'assignees':data.assigned_employee_ids})
    db.session.commit()
    return task_json(task),201


@api.get('/tasks/<int:id>')
@authenticated
def task_detail(id):
    return task_json(task_or_404(id),True)


@api.patch('/tasks/<int:id>')
@authenticated
def update_task(id):
    task = task_or_404(id)
    editable_task(task)
    merged = {key:getattr(task,key) for key in TaskInput.model_fields if hasattr(task,key)}
    merged['assigned_employee_ids'] = [a.user_id for a in task.assignments]
    data = parse(TaskInput,merged|(request.get_json() or {}))
    before = task_json(task)
    apply_task(task,data)
    set_assignees(task,data.assigned_employee_ids)
    audit('task',id,'updated',{'before':before,'after':data.model_dump(mode='json')})
    db.session.commit()
    return task_json(task)


@api.post('/tasks/<int:id>/assign')
@authenticated
def assign_task(id):
    task = task_or_404(id)
    editable_task(task)
    data = request.get_json() or {}
    if set(data) != {'assigned_employee_ids'}:
        abort(422, description='Provide assigned_employee_ids.')
    return update_task(id)


@api.post('/tasks/<int:id>/duplicate')
@authenticated
def duplicate_task(id):
    original = task_or_404(id)
    require_manager()
    task = Task(creator_id=g.user.id)
    values = {key:getattr(original,key) for key in TaskInput.model_fields if hasattr(original,key)}
    values['title'] = (original.title[:190]+' (copy)')
    values['assigned_employee_ids'] = [a.user_id for a in original.assignments]
    data = parse(TaskInput,values)
    apply_task(task,data)
    db.session.add(task)
    db.session.flush()
    set_assignees(task,data.assigned_employee_ids)
    audit('task',task.id,'duplicated',{'source':id})
    db.session.commit()
    return task_json(task),201


@api.post('/tasks/<int:id>/archive')
@authenticated
def archive_task(id):
    task = task_or_404(id)
    editable_task(task)
    task.archived = True
    audit('task',id,'archived')
    db.session.commit()
    return task_json(task)


@api.patch('/tasks/<int:id>/status')
@authenticated
def task_status(id):
    task = task_or_404(id)
    if task.archived:
        abort(409, description='Task is archived.')
    data = parse(StatusInput)
    transitions = {'Assigned':['Accepted','Cancelled'],'Accepted':['In Progress','Cancelled'],
                  'In Progress':['Pending Review','Cancelled'],'Pending Review':['Completed','Revision Requested','Cancelled'],
                  'Revision Requested':['In Progress','Pending Review','Cancelled'],'Completed':[],'Cancelled':[]}
    employee = g.user.role in STAFF
    if employee and data.status in ('Completed','Cancelled','Revision Requested'):
        abort(403, description='Only a manager can review or cancel tasks.')
    if not employee:
        editable_task(task)
    if data.status != task.status and data.status not in transitions[task.status]:
        abort(409, description=f'Invalid transition from {task.status} to {data.status}.')
    old = task.status
    task.status = data.status
    if data.progress is not None:
        if task.status in ('Completed','Cancelled') and data.status == old:
            abort(409, description='Closed task progress is locked.')
        assignment = next((a for a in task.assignments if a.user_id==g.user.id),None)
        if not assignment:
            abort(403, description='Only an assignee can update their progress.')
        assignment.progress = data.progress
    if data.status == 'Completed':
        for assignment in task.assignments:
            assignment.progress = 100
    audit('task',id,'status_changed',{'from':old,'to':data.status,'progress':data.progress})
    if employee:
        notify(task.creator_id,f'{task.title}: {data.status}',f'/tasks/{id}')
    else:
        for a in task.assignments:
            notify(a.user_id,f'{task.title}: {data.status}',f'/tasks/{id}')
    db.session.commit()
    return task_json(task)


@api.get('/tasks/<int:id>/history')
@authenticated
def history(id):
    task_or_404(id)
    return {'items':[{'id':a.id,'action':a.action,'actor_id':a.actor_id,'created_at':iso(a.created_at),
                      'details':a.details if g.user.role not in STAFF else {k:v for k,v in a.details.items() if k in ('from','to','progress')}}
                     for a in AuditLog.query.filter_by(entity='task',entity_id=id).order_by(AuditLog.id.desc())]}


@api.post('/tasks/<int:id>/comments')
@authenticated
def comment(id):
    task = task_or_404(id)
    body = (request.get_json() or {}).get('body','')
    if not isinstance(body,str) or not body.strip() or len(body)>10000:
        abort(422, description='Comment must contain 1–10000 characters.')
    db.session.add(TaskComment(task_id=id,user_id=g.user.id,body=body.strip()))
    for uid in {task.creator_id}|{a.user_id for a in task.assignments}:
        if uid != g.user.id:
            notify(uid,f'{g.user.full_name} commented on {task.title}',f'/tasks/{id}')
    audit('task',id,'comment_added')
    db.session.commit()
    return {'message':'Comment added'},201


def report_assignment(task_id, allow_closed=False):
    task = task_or_404(task_id)
    if not allow_closed and (task.archived or task.status in ('Cancelled','Completed')):
        abort(409, description='Reports cannot be changed for closed or archived tasks.')
    if not any(a.user_id==g.user.id for a in task.assignments):
        abort(403, description='You must be assigned to the related task.')
    return task


def editable_report(report):
    if report.user_id != g.user.id:
        abort(403, description='Only the author can edit the report.')
    if report.status not in ('Draft','Revision Requested'):
        abort(409, description='Only draft or revision requested reports can be edited.')


@api.get('/reports')
@authenticated
def reports():
    return paginate(report_filters(scoped_reports()),report_json,{'newest':DailyReport.id.desc(),'date':DailyReport.report_date.desc(),'oldest':DailyReport.id.asc()})


@api.post('/reports')
@authenticated
def create_report():
    data = parse(ReportInput)
    report_assignment(data.task_id)
    report = DailyReport(user_id=g.user.id)
    apply_report(report,data)
    db.session.add(report)
    db.session.flush()
    review_record(report,'Draft created')
    db.session.commit()
    return report_json(report,True),201


@api.get('/reports/<int:id>')
@authenticated
def report_detail(id):
    return report_json(report_or_404(id),True)


@api.patch('/reports/<int:id>')
@authenticated
def update_report(id):
    report = report_or_404(id)
    editable_report(report)
    merged = {key:getattr(report,key) for key in ReportInput.model_fields}
    data = parse(ReportInput,merged|(request.get_json() or {}))
    if data.task_id != report.task_id:
        abort(409, description='The related task cannot change after creation.')
    report_assignment(data.task_id,allow_closed=report.status=='Revision Requested')
    review_record(report,'Before edit')
    apply_report(report,data)
    audit('report',id,'updated')
    db.session.commit()
    return report_json(report,True)


@api.delete('/reports/<int:id>')
@authenticated
def delete_report(id):
    report = report_or_404(id)
    # Authors may discard their own drafts; administrators may remove any report.
    if not privileged() and not (report.user_id == g.user.id and report.status == 'Draft'):
        abort(403, description='Only administrators can delete submitted reports. Authors can delete their own drafts.')
    files = [item.storage_name for item in report.attachments]
    for item in list(report.attachments) + list(report.reviews):
        db.session.delete(item)
    Notification.query.filter_by(link=f'/reports/{id}').delete(synchronize_session=False)
    audit('report', id, 'deleted', {'task_id': report.task_id, 'author_id': report.user_id,
                                     'report_date': report.report_date.isoformat(), 'status': report.status})
    db.session.delete(report)
    db.session.commit()
    for name in files:
        (Path(current_app.config['UPLOAD_FOLDER'])/name).unlink(missing_ok=True)
    return {'message': 'Report deleted'}


@api.post('/reports/<int:id>/submit')
@authenticated
def submit_report(id):
    report = report_or_404(id)
    editable_report(report)
    task = report_assignment(report.task_id,allow_closed=report.status=='Revision Requested')
    report.status = 'Submitted'
    report.submitted_at = now()
    review_record(report,'Submitted')
    notify(task.creator_id,f'{g.user.full_name} submitted a report for {task.title}',f'/reports/{id}')
    db.session.commit()
    return report_json(report,True)


def reviewable(id):
    require_manager()
    report = report_or_404(id)
    # Serialize competing review actions on PostgreSQL.
    report = DailyReport.query.filter_by(id=report.id).with_for_update().populate_existing().one()
    if report.user_id == g.user.id:
        abort(403, description='You cannot review your own report.')
    if report.status not in ('Submitted','Under Review'):
        abort(409, description='Only submitted reports can be reviewed.')
    return report


@api.post('/reports/<int:id>/start-review')
@authenticated
def start_review(id):
    report = reviewable(id)
    report.status = 'Under Review'
    review_record(report,'Under Review')
    db.session.commit()
    return report_json(report,True)


@api.post('/reports/<int:id>/approve')
@authenticated
def approve_report(id):
    report = reviewable(id)
    data = parse(Feedback)
    report.status = 'Approved'
    report.manager_feedback = data.feedback
    review_record(report,'Approved',data.feedback)
    notify(report.user_id,f'Report approved: {report.task.title}',f'/reports/{id}')
    db.session.commit()
    return report_json(report,True)


@api.post('/reports/<int:id>/request-revision')
@authenticated
def request_revision(id):
    report = reviewable(id)
    data = parse(Feedback)
    if not data.feedback:
        abort(422, description='Feedback is required when requesting a revision.')
    report.status = 'Revision Requested'
    report.manager_feedback = data.feedback
    review_record(report,'Revision Requested',data.feedback)
    notify(report.user_id,f'Revision requested: {report.task.title}',f'/reports/{id}')
    db.session.commit()
    return report_json(report,True)


@api.post('/reports/<int:id>/reopen')
@authenticated
def reopen_report(id):
    require_super()
    report = report_or_404(id)
    if report.status != 'Approved':
        abort(409, description='Only approved reports can be reopened.')
    data = parse(Feedback)
    if not data.feedback:
        abort(422, description='A reason is required to reopen a report.')
    review_record(report,'Reopened',data.feedback)
    report.status = 'Revision Requested'
    report.manager_feedback = data.feedback
    notify(report.user_id,f'Report reopened: {report.task.title}',f'/reports/{id}')
    db.session.commit()
    return report_json(report,True)


def validate_upload(file):
    filename = secure_filename(file.filename or '')
    extension = Path(filename).suffix.lower()
    raw = file.read(current_app.config['MAX_CONTENT_LENGTH']+1)
    if not raw or len(raw)>current_app.config['MAX_CONTENT_LENGTH']:
        abort(422, description='File is empty or too large.')
    mime = None
    if extension == '.pdf' and raw.startswith(b'%PDF-'):
        mime = 'application/pdf'
    elif extension in ('.png','.jpg','.jpeg'):
        from PIL import Image
        try:
            im = Image.open(io.BytesIO(raw))
            if im.format != ('PNG' if extension=='.png' else 'JPEG'):
                raise ValueError()
            im.verify()
            mime = 'image/png' if extension=='.png' else 'image/jpeg'
        except Exception:
            abort(422, description='Invalid image.')
    elif extension in ('.docx','.xlsx'):
        try:
            with zipfile.ZipFile(io.BytesIO(raw)) as z:
                names = z.namelist()
                expected = 'word/document.xml' if extension=='.docx' else 'xl/workbook.xml'
                if expected not in names or '[Content_Types].xml' not in names or any('vbaproject' in n.lower() for n in names):
                    raise ValueError()
            mime = 'application/vnd.openxmlformats-officedocument.'+('wordprocessingml.document' if extension=='.docx' else 'spreadsheetml.sheet')
        except Exception:
            abort(422, description='Invalid Office file or macros detected.')
    if not mime:
        abort(422, description='Allowed file types: PDF, PNG, JPEG, DOCX and XLSX. File contents must match the extension.')
    return filename,extension,raw,mime


@api.post('/reports/<int:id>/attachments')
@api.post('/tasks/<int:id>/attachments')
@authenticated
def upload(id):
    report = None
    task = None
    if '/reports/' in request.path:
        report = report_or_404(id)
        editable_report(report)
        report_assignment(report.task_id,allow_closed=report.status=='Revision Requested')
    else:
        task = task_or_404(id)
        editable_task(task)
    file = request.files.get('file')
    if not file:
        abort(422, description='Select a file.')
    filename,extension,raw,mime = validate_upload(file)
    # Stored in the database so uploads survive serverless hosts with no persistent disk.
    attachment = Attachment(report_id=report.id if report else None,task_id=task.id if task else None,
                            owner_id=g.user.id,filename=filename,storage_name=secrets.token_hex(24)+extension,
                            mime=mime,size=len(raw),data=raw)
    db.session.add(attachment)
    db.session.flush()
    audit('report' if report else 'task',id,'attachment_added',{'filename':filename})
    db.session.commit()
    return attachment_json(attachment),201


def send_attachment(item, as_attachment):
    # Older uploads were written to UPLOAD_FOLDER; newer ones live in the database.
    source = io.BytesIO(item.data) if item.data is not None else Path(current_app.config['UPLOAD_FOLDER'])/item.storage_name
    return send_file(source,mimetype=item.mime,as_attachment=as_attachment,download_name=item.filename,max_age=0)


@api.get('/attachments/<int:id>')
@authenticated
def download(id):
    item = db.get_or_404(Attachment,id)
    if item.report_id:
        report_or_404(item.report_id)
    elif item.task_id:
        task_or_404(item.task_id)
    elif item.owner_id != g.user.id and not (g.user.role == 'super_admin' and PortfolioItem.query.filter_by(attachment_id=item.id).first()):
        abort(404)
    inline = request.args.get('inline') == '1' and item.mime.startswith('image/')
    return send_attachment(item, as_attachment=not inline)


@api.post('/auth/profile-image')
@authenticated
def profile_image():
    file = request.files.get('file')
    if not file:
        abort(422, description='Select an image.')
    filename,extension,raw,mime = validate_upload(file)
    if extension not in ('.png','.jpg','.jpeg'):
        abort(422, description='Profile images must be PNG or JPEG.')
    item = Attachment(owner_id=g.user.id,filename=filename,storage_name=secrets.token_hex(24)+extension,
                      mime=mime,size=len(raw),data=raw)
    db.session.add(item)
    db.session.flush()
    g.user.profile_image_id = str(item.id)
    audit('user',g.user.id,'profile_image_updated')
    db.session.commit()
    return user_json(g.user)


@api.get('/auth/profile-image')
@authenticated
def current_profile_image():
    if not g.user.profile_image_id:
        abort(404)
    item = Attachment.query.filter_by(id=int(g.user.profile_image_id),owner_id=g.user.id).first_or_404()
    return send_attachment(item, as_attachment=False)


@api.get('/dashboard/<role>')
@authenticated
def dashboard(role):
    if role not in ('employee','manager'):
        abort(404)
    if role == 'manager':
        require_manager()
    return analytics()


@api.get('/analytics/performance')
@api.get('/analytics/task-summary')
@authenticated
def performance():
    return analytics()


@api.get('/notifications')
@authenticated
def notifications():
    query = Notification.query.filter_by(user_id=g.user.id)
    if request.args.get('unread')=='true':
        query = query.filter_by(read=False)
    return paginate(query,lambda n:{'id':n.id,'message':n.message,'link':n.link,'read':n.read,'created_at':iso(n.created_at)}, {'newest':Notification.id.desc()})


@api.patch('/notifications/<int:id>/read')
@authenticated
def read_notification(id):
    item = Notification.query.filter_by(id=id,user_id=g.user.id).first_or_404()
    item.read = True
    db.session.commit()
    return {'message':'Marked as read'}


@api.route('/settings',methods=['GET','PATCH'])
@authenticated
def settings():
    if request.method == 'PATCH':
        require_super()
        data = request.get_json() or {}
        if set(data) != {'company_name'} or not isinstance(data['company_name'],str) or not 1 <= len(data['company_name'].strip()) <= 100:
            abort(422, description='Provide a company name of 1–100 characters.')
        db.session.merge(Setting(key='company_name',value=data['company_name'].strip()))
        audit('setting',0,'company_name_updated')
        db.session.commit()
    item = db.session.get(Setting,'company_name')
    return {'company_name':item.value if item else current_app.config['COMPANY_NAME'],
            'allow_overnight':current_app.config['ALLOW_OVERNIGHT'], 'max_upload_bytes':current_app.config['MAX_CONTENT_LENGTH']}


@api.get('/audit-logs')
@authenticated
def audit_logs():
    require_super()
    return paginate(AuditLog.query,lambda a:{'id':a.id,'actor_id':a.actor_id,'entity':a.entity,'entity_id':a.entity_id,'action':a.action,'details':a.details,'created_at':iso(a.created_at)}, {'newest':AuditLog.id.desc()})


@api.get('/exports/reports.<format>')
@authenticated
def export_reports(format):
    if format not in ('csv','xlsx','pdf'):
        abort(404)
    rows = report_filters(scoped_reports()).order_by(DailyReport.report_date.desc()).limit(10001).all()
    if len(rows)>10000:
        abort(422, description='Narrow the filters to at most 10,000 reports.')
    company = settings()['company_name']
    headers = ['Report ID','Date','Employee','Department','Task','Description','Expected output','Actual output','Target','Completed','Progress %','Hours','Status','Manager feedback','Facebook / work links (up to 15)']
    values = [[r.id,iso(r.report_date),r.user.full_name,r.user.department.name if r.user.department else '',r.task.title,r.work_description,r.expected_output,r.actual_output,r.quantity_target,r.quantity_completed,r.progress,r.total_hours,r.status,r.manager_feedback,r.evidence_url] for r in rows]
    metadata = [company,'Daily work reports',f"Date range: {request.args.get('date_from','All')} to {request.args.get('date_to','All')}",
                'Filters: '+str(dict(request.args)),f'Generated: {now().isoformat()}',f'Reports: {len(rows)} | Hours: {sum(r.total_hours for r in rows):.2f}']
    stream = io.BytesIO()
    if format == 'csv':
        text_stream = io.StringIO(newline='')
        writer = csv.writer(text_stream)
        # Prevent spreadsheet formulas injected through employee-controlled text.
        def safe(v):
            return "'"+v if isinstance(v,str) and v.lstrip().startswith(('=','+','-','@','\t','\r')) else v
        for line in metadata:
            writer.writerow([safe(line)])
        writer.writerow(headers)
        writer.writerows([[safe(v) for v in row] for row in values])
        stream.write(text_stream.getvalue().encode('utf-8-sig'))
        mime = 'text/csv; charset=utf-8'
    elif format == 'xlsx':
        from openpyxl import Workbook
        from openpyxl.styles import Font, PatternFill, Alignment
        workbook = Workbook()
        sheet = workbook.active
        sheet.title = 'Daily reports'
        for line in metadata:
            sheet.append([line])
        sheet.append(headers)
        for row in values:
            sheet.append(row)
        for row in sheet:
            for cell in row:
                if isinstance(cell.value,str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top',wrap_text=True)
        for cell in sheet[len(metadata)+1]:
            cell.font = Font(color='FFFFFF',bold=True)
            cell.fill = PatternFill('solid',fgColor='13273E')
        for column in sheet.columns:
            sheet.column_dimensions[column[0].column_letter].width = 24
        sheet.freeze_panes = f'A{len(metadata)+2}'
        sheet.auto_filter.ref = f'A{len(metadata)+1}:O{sheet.max_row}'
        workbook.save(stream)
        mime = 'application/vnd.openxmlformats-officedocument.spreadsheetml.sheet'
    else:
        from xml.sax.saxutils import escape
        from reportlab.lib import colors
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.lib.pagesizes import A4
        from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, KeepTogether
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        font = 'Helvetica'
        font_path = current_app.config['PDF_FONT_PATH']
        content = '\n'.join(metadata)+'\n'+str(values)
        if font_path:
            pdfmetrics.registerFont(TTFont('WorkdayUnicode',font_path))
            font = 'WorkdayUnicode'
        elif any(ord(c)>255 for c in content):
            abort(422, description='Configure PDF_FONT_PATH with a Khmer-capable TrueType font for Unicode PDF exports. CSV and Excel preserve Khmer without configuration.')
        styles = getSampleStyleSheet()
        for style in styles.byName.values():
            style.fontName = font
            style.shaping = 1
            if hasattr(style, 'leading'):
                style.leading = max(style.leading, style.fontSize*1.7)
        story = [Paragraph(escape(line),styles['Heading2' if i==0 else 'Normal']) for i,line in enumerate(metadata)]
        for row in values:
            story.append(Spacer(1,14))
            for label,value in zip(headers,row):
                story.append(Paragraph(escape(f'{label}: {value if value is not None else ""}').replace('\n','<br/>'),styles['Normal']))
        SimpleDocTemplate(stream,pagesize=A4,title='Daily work reports',author=company).build(story)
        mime = 'application/pdf'
    audit('export',0,'reports_'+format,{'filters':dict(request.args),'count':len(rows)})
    db.session.commit()
    stream.seek(0)
    return send_file(stream,mimetype=mime,as_attachment=True,download_name=f'work-reports-{today()}.{format}')


PORTFOLIO_KINDS = ('Poster','Video','Photo')


def portfolio_month():
    value = request.args.get('month') or today().strftime('%Y-%m')
    try:
        start = datetime.strptime(value, '%Y-%m').date()
    except ValueError:
        abort(422, description='Invalid month')
    end = (start + timedelta(days=32)).replace(day=1)
    return start, end


def portfolio_json(item):
    return {'id':item.id,'title':item.title,'kind':item.kind,'work_date':iso(item.work_date),'status':item.status,
            'department_id':item.department_id,'department':item.department.name if item.department else None,
            'user_id':item.user_id,'employee_name':item.user.full_name if item.user else None,
            'task_id':item.task_id,'task_title':item.task.title if item.task else None,
            'link':item.link,'description':item.description,
            'attachment':attachment_json(item.attachment) if item.attachment else None,'created_at':iso(item.created_at)}


@api.get('/portfolio')
@authenticated
def portfolio():
    require_super()
    start, end = portfolio_month()
    month = PortfolioItem.query.filter(PortfolioItem.work_date >= start, PortfolioItem.work_date < end)
    previous_start = (start - timedelta(days=1)).replace(day=1)
    rows = month.all()
    summary = {'total':len(rows),'pending':sum(i.status == 'Pending' for i in rows),
               'previous_total':PortfolioItem.query.filter(PortfolioItem.work_date >= previous_start, PortfolioItem.work_date < start).count(),
               'kinds':{kind:{'count':sum(i.kind == kind for i in rows),
                              'departments':len({i.department_id for i in rows if i.kind == kind and i.department_id})}
                        for kind in PORTFOLIO_KINDS}}
    query = month
    if request.args.get('kind'):
        query = query.filter(PortfolioItem.kind == request.args['kind'])
    for param, column in [('department', PortfolioItem.department_id), ('employee', PortfolioItem.user_id), ('status', PortfolioItem.status)]:
        if request.args.get(param):
            query = query.filter(column == filter_value(param, column))
    if request.args.get('q'):
        term = f"%{request.args['q']}%"
        query = query.filter(or_(PortfolioItem.title.ilike(term), PortfolioItem.description.ilike(term)))
    return paginate(query.order_by(PortfolioItem.work_date.desc(), PortfolioItem.id.desc()), portfolio_json) | {'summary':summary}


@api.post('/portfolio')
@authenticated
def create_portfolio_item():
    require_super()
    data = parse(PortfolioInput)
    for model, id in [(Department, data.department_id), (User, data.user_id), (Task, data.task_id)]:
        if id:
            db.get_or_404(model, id)
    item = PortfolioItem(**data.model_dump(exclude={'needs_approval'}), status='Pending' if data.needs_approval else 'Approved', creator_id=g.user.id)
    db.session.add(item)
    db.session.flush()
    audit('portfolio', item.id, 'created', {'title':item.title,'kind':item.kind})
    db.session.commit()
    return portfolio_json(item), 201


@api.post('/portfolio/<int:id>/file')
@authenticated
def portfolio_file(id):
    require_super()
    item = db.get_or_404(PortfolioItem, id)
    file = request.files.get('file')
    if not file:
        abort(422, description='Select a file.')
    filename,extension,raw,mime = validate_upload(file)
    if extension not in ('.png','.jpg','.jpeg','.pdf'):
        abort(422, description='Upload a JPG, PNG or PDF file. Add videos as a link.')
    previous = item.attachment
    item.attachment = Attachment(owner_id=g.user.id,filename=filename,storage_name=secrets.token_hex(24)+extension,
                                 mime=mime,size=len(raw),data=raw)
    if previous:
        db.session.delete(previous)
    audit('portfolio', id, 'file_added', {'filename':filename})
    db.session.commit()
    return portfolio_json(item)


@api.post('/portfolio/<int:id>/approve')
@authenticated
def approve_portfolio_item(id):
    require_super()
    item = db.get_or_404(PortfolioItem, id)
    item.status = 'Approved'
    audit('portfolio', id, 'approved')
    db.session.commit()
    return portfolio_json(item)


@api.delete('/portfolio/<int:id>')
@authenticated
def delete_portfolio_item(id):
    require_super()
    item = db.get_or_404(PortfolioItem, id)
    attachment = item.attachment
    db.session.delete(item)
    if attachment:
        db.session.delete(attachment)
    audit('portfolio', id, 'deleted', {'title':item.title})
    db.session.commit()
    return {'deleted':id}


@api.get('/exports/portfolio.pdf')
@authenticated
def export_portfolio():
    require_super()
    from xml.sax.saxutils import escape
    from reportlab.lib.pagesizes import A4
    from reportlab.lib.styles import getSampleStyleSheet
    from reportlab.lib.units import cm
    from reportlab.platypus import SimpleDocTemplate, Paragraph, Spacer, Image, KeepTogether
    from reportlab.pdfbase import pdfmetrics
    from reportlab.pdfbase.ttfonts import TTFont
    start, end = portfolio_month()
    items = PortfolioItem.query.filter(PortfolioItem.work_date >= start, PortfolioItem.work_date < end,
                                       PortfolioItem.status == 'Approved').order_by(PortfolioItem.work_date).all()
    company = settings()['company_name']
    pdfmetrics.registerFont(TTFont('WorkdayUnicode', current_app.config['PDF_FONT_PATH']))
    styles = getSampleStyleSheet()
    for style in styles.byName.values():
        style.fontName = 'WorkdayUnicode'
        style.shaping = 1
        if hasattr(style, 'leading'):
            style.leading = max(style.leading, style.fontSize*1.7)
    counts = ' | '.join(f'{kind}: {sum(i.kind == kind for i in items)}' for kind in PORTFOLIO_KINDS)
    story = [Paragraph(escape(company), styles['Heading2']), Paragraph(escape(f'Work portfolio {start:%m/%Y}'), styles['Heading3']),
             Paragraph(escape(f'Total: {len(items)} | {counts}'), styles['Normal'])]
    for item in items:
        block = [Spacer(1, 14), Paragraph(escape(f'{item.kind} · {item.title}'), styles['Heading4'])]
        details = [f'Date: {item.work_date:%d/%m/%Y}', f'Employee: {item.user.full_name}' if item.user else '',
                   f'Department: {item.department.name}' if item.department else '', f'Task: {item.task.title}' if item.task else '',
                   f'Link: {item.link}' if item.link else '', item.description]
        block += [Paragraph(escape(line).replace('\n', '<br/>'), styles['Normal']) for line in details if line]
        if item.attachment and item.attachment.mime.startswith('image/') and item.attachment.data:
            picture = Image(io.BytesIO(item.attachment.data))
            scale = min(8*cm / picture.imageWidth, 6*cm / picture.imageHeight)
            picture.drawWidth, picture.drawHeight = picture.imageWidth*scale, picture.imageHeight*scale
            picture.hAlign = 'LEFT'
            block.append(picture)
        story.append(KeepTogether(block))
    stream = io.BytesIO()
    SimpleDocTemplate(stream, pagesize=A4, title='Work portfolio', author=company).build(story)
    audit('export', 0, 'portfolio_pdf', {'month':f'{start:%Y-%m}','count':len(items)})
    db.session.commit()
    stream.seek(0)
    return send_file(stream, mimetype='application/pdf', as_attachment=True, download_name=f'work-portfolio-{start:%Y-%m}.pdf')
