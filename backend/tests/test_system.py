import io
from datetime import date, timedelta
import pytest
from app.extensions import db
from app.models import User, Task, TaskAssignment, DailyReport, ReportReview, AuditLog, Notification
from app.services import generate_deadline_notifications

BASE='/api/v1'


def test_pasted_work_links_normalize_bullets_and_missing_protocol(client,sign_in,report_data):
    sign_in(3)
    links='-https://www.facebook.com/profile.php?id=123\nwww.facebook.com/profile.php?id=456\n\u2022 https://example.com/work\n1. facebook.com/example'
    created=client.post(BASE+'/reports',json=report_data|{'evidence_url':links})
    assert created.status_code==201,created.json
    expected=['https://www.facebook.com/profile.php?id=123','https://www.facebook.com/profile.php?id=456','https://example.com/work','https://facebook.com/example']
    assert created.json['evidence_url'].splitlines()==expected
    assert client.get(BASE+f'/reports/{created.json["id"]}').json['evidence_url'].splitlines()==expected


def test_digital_report_progress_follows_quantity_not_links(client,sign_in,report_data,app):
    with app.app_context():
        db.session.get(User,3).role='DGT'
        db.session.commit()
    sign_in(3)
    created=client.post(BASE+'/reports',json=report_data)
    assert created.status_code==201,created.json
    assert created.json['progress']==80
    assert client.patch(BASE+f'/reports/{created.json["id"]}',json={'quantity_completed':12}).json['progress']==100


def test_fifteen_work_links_preserved_through_edit_review_and_exports(client,sign_in,report_data):
    sign_in(3)
    links=[f'https://www.facebook.com/example/posts/{index}' for index in range(15)]
    created=client.post(BASE+'/reports',json=report_data|{'evidence_url':' \r\n'+' \r\n\r\n'.join(links)+'\r\n'})
    assert created.status_code==201,created.json
    id=created.json['id']
    assert created.json['evidence_url'].splitlines()==links
    assert client.patch(BASE+f'/reports/{id}',json={'notes':'Links complete'}).json['evidence_url'].splitlines()==links
    submitted=client.post(BASE+f'/reports/{id}/submit',json={})
    assert submitted.json['reviews'][-1]['snapshot']['evidence_url'].splitlines()==links
    sign_in(2)
    assert client.get(BASE+f'/reports/{id}').json['evidence_url'].splitlines()==links
    import csv
    rows=list(csv.reader(io.StringIO(client.get(BASE+'/exports/reports.csv').data.decode('utf-8-sig'))))
    assert rows[-1][-1].splitlines()==links
    from openpyxl import load_workbook
    workbook=load_workbook(io.BytesIO(client.get(BASE+'/exports/reports.xlsx').data))
    assert workbook.active['O8'].value.splitlines()==links
    assert client.get(BASE+'/exports/reports.pdf').status_code==200


def test_work_links_maximum_count_each_length_and_validation(client,sign_in,report_data):
    sign_in(3)
    for links in [
        [f'https://example.com/{i}' for i in range(16)],
        ['https://example.com/good','javascript:alert(1)'],
        ['https://example.com/good','https://user:secret@example.com'],
        ['https://example.com/'+'x'*2048],
    ]:
        response=client.post(BASE+'/reports',json=report_data|{'evidence_url':'\n'.join(links)})
        assert response.status_code==422,response.json
    # More than the old column length is supported when each URL is valid.
    long_links=[f'https://example.com/{i}/'+('x'*200) for i in range(15)]
    created=client.post(BASE+'/reports',json=report_data|{'evidence_url':'\n'.join(long_links)})
    assert created.status_code==201,created.json
    assert len(created.json['evidence_url'])>2048


def test_report_work_link_roundtrip_review_and_export(client,sign_in,report_data):
    sign_in(3)
    link='https://www.facebook.com/example/posts/123'
    created=client.post(BASE+'/reports',json=report_data|{'evidence_url':link})
    assert created.status_code==201,created.json
    id=created.json['id']
    assert created.json['evidence_url']==link
    assert client.get(BASE+f'/reports/{id}').json['evidence_url']==link
    assert link.encode() in client.get(BASE+'/exports/reports.csv').data
    client.post(BASE+f'/reports/{id}/submit',json={})
    sign_in(2)
    reviewed=client.get(BASE+f'/reports/{id}').json
    assert reviewed['evidence_url']==link
    assert reviewed['reviews'][-1]['snapshot']['evidence_url']==link
    sign_in(4)
    assert client.get(BASE+f'/reports/{id}').status_code==404


@pytest.mark.parametrize('link',['javascript:alert(1)','data:text/html,test','not-a-link','https://user:password@example.com'])
def test_report_rejects_invalid_work_link(client,sign_in,report_data,link):
    sign_in(3)
    assert client.post(BASE+'/reports',json=report_data|{'evidence_url':link}).status_code==422


def test_login_password_and_session_revocation(client,sign_in,app):
    assert client.get(BASE+'/auth/me').status_code == 401
    assert client.post(BASE+'/auth/login',json={'username':'user3','password':'wrong'}).status_code == 401
    sign_in(3)
    assert client.get(BASE+'/auth/me').json['id']==3
    cookie=client.get_cookie('session').value
    assert client.post(BASE+'/auth/logout').status_code==200
    client.set_cookie('session',cookie)
    assert client.get(BASE+'/auth/me').status_code==401
    with app.app_context():
        assert db.session.get(User,3).password_hash!='Test-password-123'


def test_deactivated_user_blocked(client,sign_in,app):
    sign_in(3)
    with app.app_context():
        db.session.get(User,3).active=False
        db.session.commit()
    assert client.get(BASE+'/tasks').status_code==401
    assert client.post(BASE+'/auth/login',json={'username':'user3','password':'Test-password-123'}).status_code==401


@pytest.mark.parametrize('role',[2,3])
def test_record_isolation(client,sign_in,app,report_data,role):
    sign_in(4)
    report=client.post(BASE+'/reports',json=report_data|{'task_id':2}).json
    sign_in(role)
    assert client.get(BASE+'/tasks/2').status_code==404
    assert client.patch(BASE+'/tasks/2',json={'title':'Attack'}).status_code==404
    assert client.get(BASE+f'/reports/{report["id"]}').status_code==404
    assert client.patch(BASE+f'/reports/{report["id"]}',json={'notes':'Attack'}).status_code==404
    assert [t['id'] for t in client.get(BASE+'/tasks').json['items']]==[1]
    assert client.get(BASE+'/employees/4').status_code==404
    assert client.get(BASE+'/exports/reports.csv').data.count(b'Person 4')==0


def test_internal_notes_hidden_and_manager_scope(client,sign_in):
    sign_in(3)
    assert 'internal_notes' not in client.get(BASE+'/tasks/1').json
    assert client.post(BASE+'/tasks',json={}).status_code==403
    assert client.get(BASE+'/dashboard/manager').status_code==403
    sign_in(2)
    assert client.get(BASE+'/tasks/1').json['internal_notes']=='Private manager note'
    assert client.post(BASE+'/employees',json={}).status_code==422
    assert client.patch(BASE+'/departments/1',json={'name':'Other'}).status_code==403


def test_task_assignment_and_transitions(client,sign_in,app):
    sign_in(2)
    payload={'title':'New work','assigned_employee_ids':[3],'start_date':date.today().isoformat(),'deadline':date.today().isoformat()}
    response=client.post(BASE+'/tasks',json=payload)
    assert response.status_code==201,response.json
    id=response.json['id']
    assert client.post(BASE+'/tasks',json=payload|{'assigned_employee_ids':[4]}).status_code==403
    assert client.post(BASE+'/tasks',json=payload|{'deadline':'2000-01-01'}).status_code==422
    assert client.post(BASE+'/tasks',json=payload|{'priority':'Anything'}).status_code==422
    sign_in(3)
    assert client.patch(BASE+f'/tasks/{id}/status',json={'status':'Completed'}).status_code==403
    assert client.patch(BASE+f'/tasks/{id}/status',json={'status':'In Progress'}).status_code==409
    for status in ['Accepted','In Progress','Pending Review']:
        assert client.patch(BASE+f'/tasks/{id}/status',json={'status':status,'progress':80}).status_code==200
    sign_in(2)
    assert client.patch(BASE+f'/tasks/{id}/status',json={'status':'Completed'}).json['progress']==100
    assert client.patch(BASE+f'/tasks/{id}/status',json={'status':'Accepted'}).status_code==409
    assert len(client.get(BASE+f'/tasks/{id}/history').json['items'])>=5
    with app.app_context():
        assert Notification.query.filter_by(user_id=3).count()>0


def test_multi_assignment_and_archive(client,sign_in):
    sign_in(1)
    assert client.post(BASE+'/tasks/1/assign',json={'assigned_employee_ids':[3,4]}).status_code==200
    duplicate=client.post(BASE+'/tasks/1/duplicate',json={})
    assert duplicate.status_code==201
    assert len(duplicate.json['assignments'])==2
    assert client.post(BASE+'/tasks/1/archive',json={}).status_code==200
    assert 1 not in [t['id'] for t in client.get(BASE+'/tasks').json['items']]
    assert 1 in [t['id'] for t in client.get(BASE+'/tasks?archived=true').json['items']]
    assert client.get(BASE+'/tasks/1/history').status_code==200
    sign_in(3)
    assert client.patch(BASE+'/tasks/1/status',json={'status':'Accepted'}).status_code==409


def test_draft_revision_approval_and_reopen(client,sign_in,report_data,app):
    sign_in(3)
    response=client.post(BASE+'/reports',json=report_data)
    assert response.status_code==201,response.json
    r=response.json
    assert r['status']=='Draft' and r['progress']==0 and r['total_hours']==2.5  # no work links yet
    id=r['id']
    assert client.post(BASE+f'/reports/{id}/submit',json={}).json['status']=='Submitted'
    assert client.patch(BASE+f'/reports/{id}',json={'notes':'silent change'}).status_code==409
    sign_in(2)
    assert client.post(BASE+f'/reports/{id}/start-review',json={}).json['status']=='Under Review'
    assert client.post(BASE+f'/reports/{id}/request-revision',json={}).status_code==422
    assert client.post(BASE+f'/reports/{id}/request-revision',json={'feedback':'Attach supporting figures'}).json['status']=='Revision Requested'
    sign_in(3)
    assert client.patch(BASE+f'/reports/{id}',json={'evidence_url':'\n'.join(f'https://example.com/post/{i}' for i in range(15)),'notes':'Added links'}).json['progress']==100
    assert client.post(BASE+f'/reports/{id}/submit',json={}).status_code==200
    sign_in(2)
    assert client.post(BASE+f'/reports/{id}/approve',json={'feedback':'Verified'}).json['status']=='Approved'
    assert client.post(BASE+f'/reports/{id}/approve',json={}).status_code==409
    assert client.post(BASE+f'/reports/{id}/reopen',json={'feedback':'Correction'}).status_code==403
    sign_in(3)
    assert client.patch(BASE+f'/reports/{id}',json={'notes':'Attack'}).status_code==409
    sign_in(1)
    reopened=client.post(BASE+f'/reports/{id}/reopen',json={'feedback':'Correct figures'})
    assert reopened.json['status']=='Revision Requested'
    with app.app_context():
        assert ReportReview.query.filter_by(report_id=id).count()>=7
        assert AuditLog.query.filter_by(entity='report',entity_id=id,action='Reopened').count()==1


@pytest.mark.parametrize('override',[
 {'quantity_target':0},{'quantity_completed':-1},{'progress':101},{'quantity_target':'NaN'},
 {'start_time':'25:00'},{'start_time':'10:00','end_time':'10:00'},
 {'start_time':'23:00','end_time':'01:00'}, {'report_date':'2099-01-01'}, {'task_id':2}, {'user_id':4},
])
def test_report_validation(client,sign_in,report_data,override):
    sign_in(3)
    assert client.post(BASE+'/reports',json=report_data|override).status_code in (404,422)


def test_manual_progress_and_empty_submission(client,sign_in,report_data,app):
    sign_in(3)
    report=client.post(BASE+'/reports',json=report_data|{'quantity_target':None,'progress':35,'actual_output':''}).json
    assert report['progress']==0  # progress comes from work links, not a manual value
    # Description and output fields were removed from the form; an empty one no longer blocks submission.
    assert client.post(BASE+f'/reports/{report["id"]}/submit',json={}).status_code==200
    app.config['ALLOW_OVERNIGHT']=True
    overnight=client.post(BASE+'/reports',json=report_data|{'start_time':'23:30','end_time':'01:00'})
    assert overnight.json['total_hours']==1.5


def test_self_approval_forbidden(client,sign_in,report_data,app):
    with app.app_context():
        db.session.add(TaskAssignment(task_id=1,user_id=2));db.session.commit()
    sign_in(2)
    id=client.post(BASE+'/reports',json=report_data).json['id']
    client.post(BASE+f'/reports/{id}/submit',json={})
    assert client.post(BASE+f'/reports/{id}/approve',json={}).status_code==403


def test_upload_validation_and_access(client,sign_in,report_data):
    sign_in(3)
    id=client.post(BASE+'/reports',json=report_data).json['id']
    path=BASE+f'/reports/{id}/attachments'
    assert client.post(path,data={'file':(io.BytesIO(b'MZexecutable'),'evidence.exe')}).status_code==422
    assert client.post(path,data={'file':(io.BytesIO(b'MZexecutable'),'evidence.pdf')}).status_code==422
    result=client.post(path,data={'file':(io.BytesIO(b'%PDF-1.4\n%%EOF'),'evidence.pdf')})
    assert result.status_code==201,result.json
    url=result.json['url']
    assert client.get(url).status_code==200
    sign_in(4)
    assert client.get(url).status_code==404
    sign_in(2)
    assert client.get(url).status_code==200


def test_pagination_search_filters_and_exports(client,sign_in,report_data):
    sign_in(3)
    id=client.post(BASE+'/reports',json=report_data).json['id']
    client.post(BASE+f'/reports/{id}/submit',json={})
    sign_in(1)
    assert client.get(BASE+'/tasks?per_page=1').json['pages']==2
    assert client.get(BASE+'/tasks?q=Task+2').json['items'][0]['id']==2
    assert client.get(BASE+'/tasks?page=bad').status_code==422
    assert client.get(BASE+'/tasks?sort=sql-injection').status_code==422
    assert client.get(BASE+'/reports?date_from=bad').status_code==422
    assert client.get(BASE+'/reports?employee=4').json['total']==0
    csv=client.get(BASE+'/exports/reports.csv?employee=3')
    assert csv.status_code==200 and b'Person 3' in csv.data
    xlsx=client.get(BASE+'/exports/reports.xlsx')
    assert xlsx.status_code==200 and xlsx.data.startswith(b'PK')
    from openpyxl import load_workbook
    workbook=load_workbook(io.BytesIO(xlsx.data))
    assert workbook.active.max_row==8
    pdf=client.get(BASE+'/exports/reports.pdf')
    assert pdf.status_code==200 and pdf.data.startswith(b'%PDF-')
    assert b'Person 3' not in client.get(BASE+'/exports/reports.csv?employee=4').data


def test_formula_injection_export(client,sign_in,report_data):
    sign_in(3)
    client.post(BASE+'/reports',json=report_data|{'actual_output':'=HYPERLINK("evil")'})
    csv=client.get(BASE+'/exports/reports.csv')
    assert b"'=HYPERLINK" in csv.data
    xlsx=client.get(BASE+'/exports/reports.xlsx')
    from openpyxl import load_workbook
    workbook=load_workbook(io.BytesIO(xlsx.data))
    assert workbook.active['H8'].data_type=='s'


def test_metrics_cancelled_and_notifications(client,sign_in,report_data,app):
    with app.app_context():
        task=db.session.get(Task,1);task.deadline=date.today()-timedelta(days=1)
        db.session.get(Task,2).status='Cancelled';db.session.commit()
        assert generate_deadline_notifications()==1
        assert generate_deadline_notifications()==0
    sign_in(3)
    report=client.post(BASE+'/reports',json=report_data).json
    client.post(BASE+f'/reports/{report["id"]}/submit',json={})
    metrics=client.get(BASE+'/dashboard/employee').json
    assert metrics['overdue']==1 and metrics['hours']==2.5 and metrics['awaiting_review']==1
    notice=client.get(BASE+'/notifications').json['items'][0]
    assert client.patch(BASE+f'/notifications/{notice["id"]}/read',json={}).status_code==200
    sign_in(4)
    assert client.patch(BASE+f'/notifications/{notice["id"]}/read',json={}).status_code==404
    sign_in(1)
    metrics=client.get(BASE+'/analytics/performance').json
    assert metrics['employees'][3]['completion_rate']==0
    assert client.get(BASE+'/dashboard/manager?department=999').json['total_tasks']==0


def test_employee_management(client,sign_in):
    sign_in(1)
    payload={'employee_id':'NEW-1','full_name':'New Employee','username':'new.user','email':'new@example.com','password':'A-secure-password-2026','department_id':1,'manager_id':2}
    created=client.post(BASE+'/employees',json=payload)
    assert created.status_code==201,created.json
    assert 'password_hash' not in created.json
    assert client.post(BASE+'/employees',json=payload).status_code==409
    assert client.patch(BASE+'/employees/1/status',json={'active':False}).status_code==409
    assert client.patch(BASE+f'/employees/{created.json["id"]}/status',json={'active':False}).json['active']==False
    assert client.patch(BASE+'/employees/3',json={'manager_id':3}).status_code==422


def test_csrf_and_secure_headers(app):
    app.config['WTF_CSRF_ENABLED']=True
    client=app.test_client()
    assert client.post(BASE+'/auth/login',json={'username':'user3','password':'Test-password-123'}).status_code==400
    token=client.get(BASE+'/auth/csrf').json['csrf_token']
    response=client.post(BASE+'/auth/login',json={'username':'user3','password':'Test-password-123'},headers={'X-CSRFToken':token})
    assert response.status_code==200
    assert 'HttpOnly' in response.headers['Set-Cookie'] and 'SameSite=Lax' in response.headers['Set-Cookie']
    assert response.headers['X-Content-Type-Options']=='nosniff'
    assert response.headers['Cache-Control']=='no-store'


def test_password_reset_single_use_and_password_change(client,sign_in,monkeypatch,app):
    tokens=[]
    monkeypatch.setattr('app.routes.send_reset',lambda email,token:tokens.append(token))
    response=client.post(BASE+'/auth/forgot-password',json={'email':'u3@example.com'})
    assert response.status_code==200 and len(tokens)==1
    assert client.post(BASE+'/auth/forgot-password',json={'email':'absent@example.com'}).json==response.json
    reset={'token':tokens[0],'password':'Replacement-password-123'}
    assert client.post(BASE+'/auth/reset-password',json=reset).status_code==200
    assert client.post(BASE+'/auth/reset-password',json=reset).status_code==422
    assert client.post(BASE+'/auth/login',json={'username':'user3','password':reset['password']}).status_code==200
    assert client.post(BASE+'/auth/change-password',json={'current_password':reset['password'],'new_password':'Changed-password-123'}).status_code==200


def test_login_rate_limit(app):
    from app import create_app
    limited=create_app({'TESTING':True,'SQLALCHEMY_DATABASE_URI':'sqlite://','SECRET_KEY':'test','WTF_CSRF_ENABLED':False,'RATELIMIT_ENABLED':True})
    with limited.app_context():
        db.create_all()
    client=limited.test_client()
    statuses=[client.post(BASE+'/auth/login',json={'username':'absent','password':'wrong'}).status_code for _ in range(11)]
    assert statuses[-1]==429
